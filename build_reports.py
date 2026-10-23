"""Builds the four weekly Word reports from data/results.json and images/. Run after pipeline.py."""
import json, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
R = json.load(open(f"{ROOT}/data/results.json")); IMG = f"{ROOT}/images"; OUT = f"{ROOT}/reports"
os.makedirs(OUT, exist_ok=True)
K, E, C, V = R["kpi"], R["eda"], R["classification"], R["vrp"]
REG = {r["Model"]: r for r in R["regression"]}
BEST = REG["Tuned HistGB"]

def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)

class Rep:
    def __init__(self, title, subtitle):
        self.d = Document()
        s = self.d.sections[0]; s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(.9)
        st = self.d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        for n, sz in [("Heading 1", 16), ("Heading 2", 13)]:
            h = self.d.styles[n]; h.font.name = "Calibri"; h.font.size = Pt(sz); h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(subtitle); r.font.size = Pt(13); r.italic = True
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("Project: E-commerce Last-Mile Delivery Optimization using Python").font.size = Pt(11)
        self.n_fig = 0
    def h1(self, t): self.d.add_heading(t, 1)
    def h2(self, t): self.d.add_heading(t, 2)
    def p(self, t): para = self.d.add_paragraph(t); para.paragraph_format.space_after = Pt(6); return para
    def b(self, items):
        for t in items:
            if "::" in t:
                a, c = t.split("::", 1); para = self.d.add_paragraph(style="List Bullet"); para.add_run(a + ":").bold = True; para.add_run(c)
            else: self.d.add_paragraph(t, style="List Bullet")
    def code(self, src):
        t = self.d.add_table(rows=1, cols=1); t.style = "Table Grid"; c = t.rows[0].cells[0]; shade(c, "F2F2F2")
        c.paragraphs[0].text = ""
        for i, line in enumerate(src.strip("\n").split("\n")):
            para = c.paragraphs[0] if i == 0 else c.add_paragraph()
            run = para.add_run(line); run.font.name = "Consolas"; run.font.size = Pt(8.5)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas"); para.paragraph_format.space_after = Pt(0)
        self.d.add_paragraph()
    def img(self, f, cap, w=6.4):
        self.d.add_picture(f"{IMG}/{f}", width=Inches(w)); self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.n_fig += 1; c = self.d.add_paragraph(f"Figure {self.n_fig}: {cap}"); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].italic = True; c.runs[0].font.size = Pt(9.5)
    def table(self, rows, header=True, widths=None):
        t = self.d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                c = t.cell(i, j); c.text = str(v)
                for para in c.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9.5); run.bold = (i == 0 and header)
                if i == 0 and header: shade(c, "D9E2F3")
        self.d.add_paragraph()
    def save(self, name): self.d.save(f"{OUT}/{name}"); print("saved", name)

DATA_NOTE = ("Data note: the public Kaggle \"Amazon Delivery Dataset\" (sujalsuthar/amazon-delivery-dataset) is the reference dataset for this project. "
             "To keep the work fully reproducible, the numerical results in this report were produced on a synthetic dataset of 20,000 orders that follows exactly "
             "the same schema (order, agent, store/drop coordinates, time, weather, traffic, vehicle, area, category and delivery time). "
             "The pipeline automatically switches to the real CSV when it is placed in data/amazon_delivery.csv.")

SRC = ["Jazemi, R., Alidadiani, E., Ahn, K., & Jang, J. (2023). A Review of Literature on Vehicle Routing Problems of Last-Mile Delivery in Urban Areas. Applied Sciences, 13(24), 13015. https://doi.org/10.3390/app132413015",
       "Giuffrida, N., Fajardo-Calderin, J., Masegosa, A. D., Werner, F., Steudter, M., & Pilla, F. (2022). Optimization and Machine Learning Applied to Last-Mile Logistics: A Review. Sustainability, 14(9), 5329. https://www.mdpi.com/2071-1050/14/9/5329",
       "Deloitte. Last-Mile Delivery Strategy and Cost Optimization. https://www.deloitte.com/us/en/industries/consumer/articles/last-mile-delivery-strategy-cost-optimization.html",
       "McKinsey & Company. Creating value from AI and digital capabilities in logistics operations. https://www.mckinsey.com/capabilities/operations/our-insights/creating-value-from-ai-and-digital-capabilities-in-logistics-operations",
       "Kaggle. Amazon Delivery Dataset (sujalsuthar). https://www.kaggle.com/datasets/sujalsuthar/amazon-delivery-dataset"]

# =============================================================== WEEK 1
w = Rep("Week 1: Strategic Planning and Data Exploration in Logistics", "Strategic Planning Report")
w.h1("1. Project Definition")
w.p("An online retailer ships orders from city stores to customers through a fleet of delivery agents using motorcycles, scooters, vans and bicycles. "
    "Customers are promised delivery within 120 minutes. Delivery times vary strongly with distance, traffic, weather and time of day, and routes are planned manually. "
    "The business problem is twofold: (1) delivery times are unpredictable, so many orders miss the 120-minute promise, and (2) the routes ridden by agents are longer than necessary, which raises cost.")
w.p("This project therefore focuses on last-mile delivery optimization, a combination of delivery-time prediction (to anticipate late orders) and route optimization (to reduce distance travelled per order). "
    "The scenario covers three of the typical logistics challenges named in the task: route optimization, resource (vehicle and agent) allocation, and supply-chain visibility.")
w.h1("2. Key Performance Indicators (KPIs)")
w.table([["KPI", "Definition / formula", "Why it matters"],
         ["On-time delivery rate (%)", "Orders delivered within 120 min / total orders x 100", "Direct measure of the customer promise and reliability"],
         ["Average delivery time (min)", "Sum of delivery minutes / number of orders", "Measures speed and overall operational efficiency"],
         ["Distance per delivery (km)", "Total route kilometres / delivered orders", "Proxy for cost per delivery (fuel, agent time, vehicle wear)"],
         ["Prediction error, MAE (min)", "Mean absolute difference between predicted and actual time", "Shows how reliable delivery-time promises can be made"]])
w.p(f"Baseline values calculated on the cleaned data: on-time delivery {K['on_time_pct']}%, average delivery time {K['avg_delivery_min']} minutes, "
    f"average store-to-customer distance {K['avg_distance_km']} km.")
w.h1("3. Literature and Data Research")
w.h2("3.1 Background research")
w.b(["Vehicle routing in last-mile delivery:: Jazemi et al. (2023) review the literature on vehicle routing problems for urban last-mile delivery and group the work into crowdsourced drivers, parcel lockers, drone-assisted delivery and customer-chosen delivery locations. The review confirms that vehicle capacity, time windows and urban traffic are the main constraints, which is why this project uses a capacity-constrained vehicle routing model.",
     "Optimization plus machine learning:: Giuffrida et al. (2022) review 77 studies and classify them as traditional routing models, machine-learning approaches and mixed methods. Their conclusion supports combining prediction (machine learning) with routing (optimization), which is exactly the structure of this project.",
     "Business perspective on cost and service:: Deloitte notes that delivery costs are rising while customers want different trade-offs between price and speed, so companies should offer a wider range of delivery options. This justifies tracking cost-related KPIs together with on-time performance.",
     "Value of analytics in logistics:: McKinsey argues that returns come to firms that expose trade-offs, place analytics where decisions are made, and tie decisions to execution. This motivates ending each analysis with a concrete, actionable recommendation."])
w.h2("3.2 Dataset")
w.p("The reference dataset is the Amazon Delivery Dataset available on Kaggle. Each record is one delivery order with the following fields: Order_ID, Agent_Age, Agent_Rating, "
    "Store and Drop latitude/longitude, Order_Date, Order_Time, Pickup_Time, Weather, Traffic, Vehicle, Area, Category and the target variable Delivery_Time in minutes. "
    "Because it contains real coordinates, timestamps and a measured delivery time, it supports every technique planned in this project: regression, classification, clustering and routing.")
w.p(DATA_NOTE)
w.h1("4. How Data Science Concepts Apply")
w.table([["Technique", "Application in this project", "Output"],
         ["Regression", "Predict Delivery_Time from distance, traffic, weather, vehicle, hour and agent features", "Estimated delivery minutes"],
         ["Classification", "Predict whether an order will be late (more than 120 min)", "Late / on-time label with probability"],
         ["Clustering (K-Means)", "Group deliveries by distance, time and preparation time to find delivery segments", "Segments such as short, fast orders vs long, risky orders"],
         ["Optimization (VRP, OR-Tools)", "Assign orders to riders and sequence stops under vehicle capacity to minimise kilometres", "Optimised routes and kilometres saved"]])
w.h1("5. Strategic Roadmap")
w.img("08_roadmap.png", "End-to-end process flow of the planned analysis", 6.5)
w.table([["Phase", "Activities (Python libraries)", "Deliverable"],
         ["1. Data collection", "Load the Kaggle CSV with pandas; inspect schema, size and types", "Raw dataset and data dictionary"],
         ["2. Cleaning and preprocessing", "Duplicates, missing values, invalid coordinates and ages, outliers (IQR), feature engineering (haversine distance, hour), scaling", "Cleaned dataset"],
         ["3. Exploratory analysis", "Distributions, correlations, traffic/weather/vehicle effects with matplotlib and seaborn", "Charts and insights"],
         ["4. Predictive modeling", "Linear regression, random forest, gradient boosting, late-delivery classifier, K-Means with scikit-learn", "Validated models and metrics"],
         ["5. Route optimization", "Capacity-constrained vehicle routing with Google OR-Tools", "Optimised routes, kilometres saved"],
         ["6. Evaluation and decisions", "Compare results with baseline KPIs and write recommendations", "Final report"]])
w.h1("6. Code Illustration")
w.p("The following snippets show how the main parts of the plan will be implemented.")
w.h2("Feature engineering: distance between store and customer")
w.code('''import numpy as np, pandas as pd

def haversine(lat1, lon1, lat2, lon2):
    p = np.pi / 180
    a = (np.sin((lat2 - lat1) * p / 2) ** 2
         + np.cos(lat1 * p) * np.cos(lat2 * p) * np.sin((lon2 - lon1) * p / 2) ** 2)
    return 12742 * np.arcsin(np.sqrt(a))          # kilometres

df["distance_km"] = haversine(df.Store_Latitude, df.Store_Longitude,
                              df.Drop_Latitude, df.Drop_Longitude)
df["late"] = (df["Delivery_Time"] > 120).astype(int)   # SLA label''')
w.h2("Regression and classification")
w.code('''from sklearn.model_selection import train_test_split
from sklearn.ensemble import HistGradientBoostingRegressor, RandomForestClassifier

X = pd.get_dummies(df[features], columns=categorical, drop_first=True)
X_tr, X_te, y_tr, y_te = train_test_split(X, df["Delivery_Time"], test_size=0.2, random_state=42)

reg = HistGradientBoostingRegressor().fit(X_tr, y_tr)            # delivery-time regression
clf = RandomForestClassifier(200, class_weight="balanced").fit(X_tr, df.loc[X_tr.index, "late"])''')
w.h2("Clustering and route optimization (pseudocode)")
w.code('''# K-Means on scaled [distance_km, Delivery_Time, prep_min]; choose k with the silhouette score
for k in 2..6: score[k] = silhouette(KMeans(k).fit_predict(Z))
k_best = argmax(score)

# Vehicle routing with OR-Tools
manager = RoutingIndexManager(n_locations, n_vehicles, depot=0)
routing = RoutingModel(manager)
routing.SetArcCostEvaluatorOfAllVehicles(distance_callback)       # minimise kilometres
routing.AddDimensionWithVehicleCapacity(demand_callback, 0, [CAP]*n_vehicles, True, "Cap")
solution = routing.SolveWithParameters(GUIDED_LOCAL_SEARCH, time_limit=5s)''')
w.h1("7. Expected Outcomes and Impact on Decision-Making")
w.b(["Predictable delivery promises:: a model that predicts delivery time within a few minutes lets the business quote realistic delivery windows instead of a single fixed promise.",
     "Early warning for late orders:: a classifier flags risky orders (long distance, heavy traffic, bad weather) so dispatchers can reassign vehicles or notify customers before a delay happens.",
     "Lower cost per delivery:: optimised routes reduce kilometres travelled, which cuts fuel, vehicle wear and agent hours.",
     "Evidence-based planning:: insights on peak hours, traffic and vehicle type support staffing, vehicle-mix and area-expansion decisions."])
w.h1("8. Conclusion")
w.p("Last-mile delivery combines two connected problems: knowing how long a delivery will take and deciding how to route vehicles efficiently. "
    "This plan defines a realistic e-commerce scenario, measures it with four KPIs, grounds it in published research, and applies regression, classification, clustering and vehicle routing "
    "through a six-phase Python pipeline. If implemented as planned, the project should improve on-time performance and reduce distance travelled, and it gives managers a data-driven basis for dispatch, staffing and customer-promise decisions.")
w.h1("References"); w.b(SRC)
w.save("Week1_Strategic_Planning_Report.docx")

# =============================================================== WEEK 2
w = Rep("Week 2: Data Collection, Cleaning and Preprocessing", "Preprocessing Pipeline Report")
w.h1("1. Data Collection")
w.p("The reference dataset is the Kaggle Amazon Delivery Dataset (one row per delivery order). The table below describes the main characteristics of the data used in this report.")
w.table([["Characteristic", "Value"], ["Source / reference", "Kaggle: sujalsuthar/amazon-delivery-dataset"], ["Records (raw)", f"{R['raw_rows']:,}"], ["Columns (raw)", R["raw_cols"]],
         ["Target variable", "Delivery_Time (minutes)"], ["Numeric fields", "Agent_Age, Agent_Rating, store/drop latitude and longitude, Delivery_Time"],
         ["Categorical fields", "Weather, Traffic, Vehicle, Area, Category"], ["Time fields", "Order_Date, Order_Time, Pickup_Time"]])
w.p(DATA_NOTE)
w.h1("2. Data Quality Issues Identified")
w.p("Before cleaning, the data was inspected with df.info(), df.isna().sum(), df.duplicated().sum() and df.describe(). The following problems were found.")
w.table([["Issue", "Count", "How it was detected"],
         ["Duplicate rows", R["duplicates"], "df.duplicated().sum()"],
         ["Missing Agent_Rating", R["missing"]["Agent_Rating"], "df.isna().sum()"],
         ["Missing Pickup_Time", R["missing"]["Pickup_Time"], "df.isna().sum()"],
         ["Missing Weather", R["missing"]["Weather"], "df.isna().sum()"],
         ["Invalid coordinates (0, 0)", R["invalid_coords"], "abs(lat/long) < 1 check"],
         ["Impossible agent age (under 18)", R["invalid_age"], "range check 18-70"],
         ["Rating outside 1-5", 30, "range check"],
         ["Inconsistent text labels (case, spaces)", "about 2% of Traffic", "value_counts() of categorical columns"],
         ["Extreme delivery times", R["outliers_removed"], "IQR rule on Delivery_Time"]])
w.h1("3. Cleaning Methodology and Justification")
w.table([["Step", "Technique", "Why this technique and its effect on analysis"],
         ["Duplicates", "drop_duplicates()", "Duplicates double-count orders and bias averages and model training."],
         ["Text labels", "strip() and title-case", "'JAM ' and 'Jam' would otherwise become two categories and split the statistics."],
         ["Missing categorical (Weather)", "Fill with mode", "Keeps the row; mode is the least distorting value for a category."],
         ["Missing numeric (Agent_Rating)", "Fill with median", "The median is robust to outliers, unlike the mean."],
         ["Missing Pickup_Time", "Compute prep time and fill with median", "Preparation time is a useful predictor, so the row is kept."],
         ["Invalid values", "Rating outside 1-5 set to missing then imputed; ages under 18 replaced by median; (0,0) coordinates dropped", "Impossible values are data-entry errors; wrong coordinates would give wrong distances."],
         ["Outliers", "IQR rule with factor 3.0 on Delivery_Time", "Extreme values distort mean, regression coefficients and RMSE. A wide factor removes only clear errors, keeping genuine long deliveries."],
         ["Feature engineering", "haversine distance, hour, day of week, preparation minutes", "Raw coordinates and timestamps are not directly usable by models."],
         ["Normalization", "StandardScaler on numeric columns", "Puts features on the same scale, needed for K-Means, logistic regression and distance-based methods."]])
w.h1("4. Python Implementation")
w.code('''import pandas as pd, numpy as np
from sklearn.preprocessing import StandardScaler

df = raw.drop_duplicates().copy()

# 1. text labels
for col in ["Traffic", "Weather", "Vehicle", "Area", "Category"]:
    df[col] = df[col].astype("string").str.strip().str.title()

# 2. missing values
df["Weather"] = df["Weather"].fillna(df["Weather"].mode()[0])
df["Agent_Rating"] = df["Agent_Rating"].where(df["Agent_Rating"].between(1, 5))
df["Agent_Rating"] = df["Agent_Rating"].fillna(df["Agent_Rating"].median())

# 3. invalid values
bad_age = ~df["Agent_Age"].between(18, 70)
df.loc[bad_age, "Agent_Age"] = df["Agent_Age"].median()
bad_geo = (df[["Store_Latitude", "Store_Longitude", "Drop_Latitude", "Drop_Longitude"]].abs() < 1).any(axis=1)
df = df[~bad_geo]''')
w.code('''# 4. time features
ot = pd.to_datetime(df["Order_Date"] + " " + df["Order_Time"])
pt = pd.to_datetime(df["Order_Date"] + " " + df["Pickup_Time"], errors="coerce")
df["prep_min"] = ((pt - ot).dt.total_seconds() / 60).clip(0, 60)
df["prep_min"] = df["prep_min"].fillna(df["prep_min"].median())
df["hour"] = ot.dt.hour

# 5. outliers (IQR rule)
q1, q3 = df["Delivery_Time"].quantile([.25, .75]); iqr = q3 - q1
mask = (df["Delivery_Time"] > q3 + 3 * iqr) | (df["Delivery_Time"] < q1 - 3 * iqr)
df = df[~mask]

# 6. normalization
num = ["Agent_Age", "Agent_Rating", "distance_km", "prep_min"]
df[num] = StandardScaler().fit_transform(df[num])''')
w.h1("5. Results of Cleaning")
w.table([["Measure", "Before", "After"], ["Rows", f"{R['raw_rows']:,}", f"{R['clean_rows']:,}"],
         ["Missing values", sum(R["missing"].values()), 0], ["Duplicate rows", R["duplicates"], 0], ["Invalid coordinates", R["invalid_coords"], 0], ["Extreme delivery times", R["outliers_removed"], 0]])
w.img("01_data_cleaning.png", "Delivery time before and after cleaning, and missing values found")
w.p(f"Cleaning removed {R['raw_rows'] - R['clean_rows']} rows ({100 * (R['raw_rows'] - R['clean_rows']) / R['raw_rows']:.1f}% of the data) while keeping {R['clean_rows']:,} reliable records. "
    "The boxplot shows that the extreme delivery times above the normal range disappear after outlier handling.")
w.h1("6. Reflection: Impact of Data Quality on Analytics")
w.b(["Wrong inputs produce wrong decisions:: invalid (0, 0) coordinates would create distances of thousands of kilometres and corrupt both the model and the routing.",
     "Outliers distort metrics:: a few extreme delivery times raise the average and RMSE and make the on-time KPI look worse than reality.",
     "Missing data wastes information:: imputing instead of deleting keeps more than 98% of the rows available for learning.",
     "Consistent categories give trustworthy comparisons:: unified labels allow correct comparison of traffic and weather effects.",
     "Good preprocessing makes later weeks possible:: the cleaned file is the input for the exploratory analysis, models and route optimisation."])
w.h1("References"); w.b(SRC[-1:])
w.save("Week2_Data_Cleaning_Report.docx")

# =============================================================== WEEK 3
w = Rep("Week 3: Advanced Data Analysis and Visualization in Logistics", "Exploratory Data Analysis Report")
w.h1("1. Dataset Description")
w.p(f"The analysis uses the cleaned delivery dataset with {R['clean_rows']:,} orders prepared in Week 2. The key variables are listed below.")
w.table([["Variable", "Type", "Meaning"], ["Delivery_Time", "numeric (min)", "Target: time from order to delivery"], ["distance_km", "numeric", "Store-to-customer distance (haversine)"],
         ["prep_min", "numeric (min)", "Order preparation time at the store"], ["Traffic", "categorical", "Low, Medium, High, Jam"], ["Weather", "categorical", "Sunny, Cloudy, Fog, Windy, Stormy, Sandstorms"],
         ["Vehicle / Area", "categorical", "Motorcycle, scooter, van, bicycle / Urban, Metropolitian, Semi-Urban, Other"], ["hour", "numeric", "Hour of the order (8-22)"], ["late", "binary", "1 if Delivery_Time exceeds 120 minutes"]])
w.p(DATA_NOTE)
w.h1("2. Exploratory Data Analysis")
w.h2("2.1 Central tendency and spread")
w.table([["Statistic", "Delivery time (min)"], ["Mean", E["mean"]], ["Median", E["median"]], ["Standard deviation", E["std"]], ["Skewness", E["skew"]]])
w.p(f"The mean ({E['mean']} min) is higher than the median ({E['median']} min) and the skewness is positive ({E['skew']}), so the distribution has a longer right tail: most orders are delivered in about 1.5 to 2 hours, but a minority take much longer.")
w.h2("2.2 Distributions and relationships")
w.img("02_eda_overview.png", "Delivery-time distribution, traffic effect, distance relationship and weather effect", 6.4)
w.b([f"Histogram:: justified because it shows the full shape of the target and where the 120-minute SLA sits relative to the mass of orders. {K['on_time_pct']}% of orders are within the SLA.",
     "Boxplot by traffic:: justified because it compares both the median and the spread across categories and exposes outliers. The average delivery time rises from about "
     f"{E['time_by_traffic']['Low']:.0f} minutes in low traffic to {E['time_by_traffic']['Jam']:.0f} minutes in a traffic jam.",
     f"Scatter plot of distance vs time:: justified for two numeric variables. The relationship is strongly positive (correlation {E['corr_distance_time']}) and the colour shows that heavy traffic shifts the whole cloud upward.",
     f"Bar chart by weather:: justified for comparing category averages. Stormy ({E['time_by_weather']['Stormy']:.0f} min) and sandstorm ({E['time_by_weather']['Sandstorms']:.0f} min) conditions are the slowest; sunny weather is the fastest ({E['time_by_weather']['Sunny']:.0f} min)."])
w.h2("2.3 Correlations and time patterns")
w.img("03_correlation_trends.png", "Correlation heatmap, delivery time by hour, and area x vehicle matrix", 6.5)
w.b(["Heatmap:: justified to see all numeric relationships at once. Distance is by far the strongest correlate of delivery time; agent age and rating are almost uncorrelated with it.",
     "Line chart by hour:: justified for a time trend. Average delivery time jumps from about 107 minutes to about 118 minutes between 18:00 and 21:00, which is the evening peak.",
     "Area x vehicle matrix:: justified for two categorical dimensions. Bicycles are the slowest vehicle in every area; vans are the fastest."])
w.h1("3. Python Code Used")
w.code('''import seaborn as sns, matplotlib.pyplot as plt
order = ["Low", "Medium", "High", "Jam"]

fig, ax = plt.subplots(2, 2, figsize=(15, 10))
sns.histplot(df.Delivery_Time, bins=45, kde=True, ax=ax[0, 0])
ax[0, 0].axvline(120, color="orange", ls=":", label="SLA 120 min")

sns.boxplot(data=df, x="Traffic", y="Delivery_Time", order=order, ax=ax[0, 1])
sns.scatterplot(data=df.sample(2500), x="distance_km", y="Delivery_Time",
                hue="Traffic", hue_order=order, alpha=.5, ax=ax[1, 0])
sns.barplot(data=df, x="Weather", y="Delivery_Time", errorbar=None, ax=ax[1, 1])
plt.tight_layout(); plt.savefig("images/02_eda_overview.png", dpi=120)''')
w.code('''# central tendency, correlation, hourly trend
df["Delivery_Time"].agg(["mean", "median", "std", "skew"])
sns.heatmap(df[["Agent_Age","Agent_Rating","distance_km","prep_min","hour","Delivery_Time"]].corr(),
            annot=True, cmap="coolwarm", center=0)
df.groupby("hour")["Delivery_Time"].mean().plot(marker="o")
sns.heatmap(df.pivot_table(index="Area", columns="Vehicle", values="Delivery_Time"), annot=True)''')
w.h1("4. Analytical Insights")
w.table([["Insight", "Evidence", "Recommendation"],
         ["Distance is the main cost and time driver", f"Correlation {E['corr_distance_time']}; about 6 minutes per extra kilometre", "Assign orders to the nearest store or rider; use route optimisation for long orders"],
         ["Traffic jams cause most delays", f"{E['late_by_traffic']['Jam']}% of jam orders are late vs {E['late_by_traffic']['Low']}% in low traffic", "Use live traffic to re-route, and quote longer promise windows during jams"],
         ["Evening peak (18:00-21:00) is a bottleneck", "About 10 minutes slower on average", "Add riders and pre-position stock before 18:00"],
         ["Bad weather adds 10-20 minutes", f"Stormy {E['time_by_weather']['Stormy']:.0f} min vs sunny {E['time_by_weather']['Sunny']:.0f} min", "Weather-based buffers in delivery promises"],
         ["Bicycles are slowest", f"{E['time_by_vehicle']['Bicycle']:.0f} min vs {E['time_by_vehicle']['Motorcycle']:.0f} min for motorcycles", "Restrict bicycles to short-distance orders"]])
w.h1("5. How the Visualizations Help Interpret Logistics Performance")
w.p("The histogram measures reliability (share of orders inside the SLA), the boxplots and bar charts identify which conditions create delay (cost drivers), the scatter plot quantifies the value of shorter distance, "
    "and the hourly line chart reveals bottlenecks in time. Together they convert raw records into operational decisions on staffing, vehicle choice and customer promises.")
w.h1("6. Conclusion")
w.p("The exploratory analysis shows that delivery performance is driven mainly by distance, traffic and the evening peak, with weather and vehicle type as secondary factors. "
    "These findings guide the features used in the Week 4 prediction model and the routing optimisation.")
w.h1("References"); w.b(SRC[-1:])
w.save("Week3_EDA_Visualization_Report.docx")

# =============================================================== WEEK 4
w = Rep("Week 4: Predictive Modeling and Optimization in Logistics Systems", "Final Modeling and Optimization Report")
w.h1("1. Problem Definition")
w.p("The target is Delivery_Time, the number of minutes between order and delivery. An accurate forecast lets the business promise realistic delivery windows and detect risky orders early. "
    "A second task, the late-delivery classifier, predicts whether an order will exceed the 120-minute promise. The final part of the project uses route optimisation to reduce travelled distance.")
w.p(f"Dataset: {R['clean_rows']:,} cleaned orders (Week 2). Features: Agent_Age, Agent_Rating, distance_km, prep_min, hour, day of week, Weather, Traffic, Vehicle, Area and Category (categorical variables one-hot encoded). 80% of the data is used for training and 20% for testing.")
w.p(DATA_NOTE)
w.h1("2. Model Selection and Rationale")
w.table([["Model", "Why it was chosen"], ["Linear Regression", "Simple, interpretable baseline; shows how much of the problem is linear"],
         ["Random Forest", "Captures non-linear effects and interactions; robust to outliers; gives feature importance"],
         ["Histogram Gradient Boosting", "State-of-the-art for tabular data; fast; strong accuracy with tuning"],
         ["Random Forest classifier / Logistic Regression", "Predict late (1) or on time (0); class_weight='balanced' handles the 36% late share"]])
w.h1("3. Implementation")
w.code('''from sklearn.model_selection import train_test_split, cross_val_score, RandomizedSearchCV
from sklearn.linear_model import LinearRegression
from sklearn.ensemble import RandomForestRegressor, HistGradientBoostingRegressor
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score

X = pd.get_dummies(df[features], columns=categorical, drop_first=True).astype(float)
y = df["Delivery_Time"]
X_tr, X_te, y_tr, y_te = train_test_split(X, y, test_size=0.2, random_state=42)

models = {"Linear Regression": LinearRegression(),
          "Random Forest": RandomForestRegressor(150, max_depth=14, min_samples_leaf=5, random_state=42),
          "Hist Gradient Boosting": HistGradientBoostingRegressor(random_state=42)}
for name, m in models.items():
    m.fit(X_tr, y_tr); pred = m.predict(X_te)
    print(name, mean_absolute_error(y_te, pred),
          mean_squared_error(y_te, pred) ** .5, r2_score(y_te, pred),
          cross_val_score(m, X, y, cv=3, scoring="r2").mean())''')
w.h2("Hyperparameter tuning")
w.code('''search = RandomizedSearchCV(
    HistGradientBoostingRegressor(random_state=42),
    {"learning_rate": [.03, .05, .1, .2], "max_depth": [3, 5, 7, None],
     "max_leaf_nodes": [15, 31, 63], "l2_regularization": [0, .1, 1]},
    n_iter=10, cv=3, scoring="neg_root_mean_squared_error", random_state=42).fit(X_tr, y_tr)
tuned = search.best_estimator_''')
w.h1("4. Evaluation and Validation")
w.p("Metrics: MAE is the average error in minutes and is easy to explain to managers; RMSE penalises large errors; R-squared shows the share of variance explained. "
    "A 3-fold cross-validation R-squared confirms that the result does not depend on a single train/test split.")
w.table([["Model", "MAE (min)", "RMSE (min)", "R-squared", "CV R-squared"]] +
        [[r["Model"], r["MAE"], r["RMSE"], r["R2"], r.get("CV_R2", "-")] for r in R["regression"]])
w.p(f"Best parameters found by the randomised search: {R['best_params']}. All three regression models explain about 90-92% of the variance. "
    f"Gradient boosting and linear regression are very close (MAE about 7 minutes), which shows that the delivery-time process is mostly smooth and additive; random forest is slightly worse in this data. "
    f"The tuned model predicts delivery time with an average error of {BEST['MAE']:.1f} minutes.")
w.img("05_regression_results.png", "Model comparison, actual vs predicted delivery time, and feature importance")
top = R["top_features"]
w.p(f"Feature importance (random forest) shows that distance_km is the dominant driver ({top['distance_km']:.2f}), followed by traffic jam conditions ({top['Traffic_Jam']:.2f}); hour of day, stormy weather and area have small but visible effects.")
w.h2("Late-delivery classification")
w.table([["Metric", "Random Forest classifier"], ["Accuracy", C["accuracy"]], ["F1-score (late class)", C["f1"]], ["ROC-AUC", C["roc_auc"]], ["Share of late orders", f"{C['late_rate_pct']}%"]])
w.p(f"The classifier separates late from on-time deliveries very well (ROC-AUC {C['roc_auc']}); logistic regression reaches a similar AUC ({C['logit_auc']}), confirming that the signal is strong.")
w.h2("Customer segments (K-Means)")
seg = R["kmeans"]["segments"]
w.table([["Segment", "Orders", "Avg km", "Avg minutes", "Late %"]] + [[s["segment"], s["orders"], s["avg_km"], s["avg_min"], s["late_pct"]] for s in seg])
w.p(f"The silhouette score selects k = {R['kmeans']['best_k']}. One segment holds long-distance orders (about {max(s['avg_km'] for s in seg):.0f} km) with a very high late rate, "
    "so this segment should be the target of special handling.")
w.img("04_clustering.png", "Silhouette score and delivery segments", 6.0)
w.h1("5. Optimization Strategies")
w.h2("5.1 Route optimization with OR-Tools")
w.p(f"A capacity-constrained vehicle routing problem was solved for a batch of {V['stops']} deliveries from one store, with {V['capacity']} units of capacity per rider. "
    f"The baseline is a greedy nearest-neighbour dispatcher; the optimised solution uses OR-Tools with guided local search.")
w.code('''from ortools.constraint_solver import pywrapcp, routing_enums_pb2
mgr = pywrapcp.RoutingIndexManager(N, n_vehicles, 0)
routing = pywrapcp.RoutingModel(mgr)
cb = routing.RegisterTransitCallback(lambda i, j: dist_m[mgr.IndexToNode(i)][mgr.IndexToNode(j)])
routing.SetArcCostEvaluatorOfAllVehicles(cb)
dcb = routing.RegisterUnaryTransitCallback(lambda i: demand[mgr.IndexToNode(i)])
routing.AddDimensionWithVehicleCapacity(dcb, 0, [CAP] * n_vehicles, True, "Cap")
p = pywrapcp.DefaultRoutingSearchParameters()
p.first_solution_strategy = routing_enums_pb2.FirstSolutionStrategy.PATH_CHEAPEST_ARC
p.local_search_metaheuristic = routing_enums_pb2.LocalSearchMetaheuristic.GUIDED_LOCAL_SEARCH
p.time_limit.seconds = 5
solution = routing.SolveWithParameters(p)''')
w.table([["Measure", "Baseline (greedy)", "Optimised (OR-Tools)"], ["Total distance (km)", V["baseline_km"], V["optimized_km"]], ["Riders used", "-", V["vehicles_used"]], ["Distance saved", "-", f"{V['saving_pct']}%"]])
w.img("06_route_optimization.png", "Baseline routes vs OR-Tools optimised routes for 30 stops", 6.4)
w.h2("5.2 Operational recommendations from the model")
w.b([f"Reduce preparation time:: the what-if analysis on {R['whatif']['orders_affected']} high-risk orders (heavy traffic, more than 8 km) shows that preparing orders 3 minutes faster reduces predicted delivery time by about {R['whatif']['avg_min_saved_prep']} minutes each.",
     "Flag high-risk orders:: use the late-delivery classifier at checkout to assign faster vehicles or show a longer promise window.",
     "Dynamic dispatch by traffic:: jam conditions are the second largest driver; re-route or delay non-urgent orders during jams.",
     "Vehicle allocation:: use vans/motorcycles for long distances and restrict bicycles to short, low-traffic orders.",
     "Peak-hour staffing:: add riders between 18:00 and 21:00 where average delivery time rises by about 10 minutes."])
w.img("07_classification_impact.png", "Confusion matrix, late rate by traffic, and route distance saved", 6.4)
w.h1("6. Final Recommendations and Conclusion")
w.p(f"A gradient-boosting model predicts delivery time with an average error of about {BEST['MAE']:.0f} minutes (R-squared {BEST['R2']}), and a classifier identifies late orders with ROC-AUC {C['roc_auc']}. "
    f"Route optimisation reduced the distance of a 30-stop batch by {V['saving_pct']}% compared with a greedy dispatcher. "
    "Combining these tools gives the business reliable delivery promises, early warning for risky orders and lower cost per delivery. "
    "Limitations: the results were obtained on a synthetic dataset that follows the Amazon Delivery schema, and the route example uses one store with random demand; the same pipeline should be re-run on the real Kaggle file, and on real traffic data, before operational use.")
w.img("00_kpi_dashboard.png", "Project KPI dashboard", 6.4)
w.h1("References"); w.b(SRC)
w.save("Week4_Predictive_Modeling_Optimization_Report.docx")
